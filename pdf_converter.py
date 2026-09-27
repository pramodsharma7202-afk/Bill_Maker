"""Render the generated invoice workbook to PDF using Python libraries only."""

from pathlib import Path
from copy import copy
import re
import tempfile

from openpyxl import load_workbook
from xhtml2pdf import pisa
from xlsx2html import xlsx2html

from billing import amount_in_words


def _prepare_pdf_copy(xlsx_path: Path, data: dict, destination: Path) -> None:
    """Copy the workbook and replace formulas with their final display values.

    openpyxl intentionally does not calculate formulas. The Excel invoice keeps
    its editable formulas; only this temporary rendering copy receives values.
    """
    workbook = load_workbook(xlsx_path)
    sheet = workbook["Tax Invoice"]

    for row in range(17, 28):
        item_index = row - 17
        item = data["items"][item_index] if item_index < len(data["items"]) else None
        cell = sheet.cell(row, 8)
        cell.value = f"Rs. {item['amount']:,.2f}" if item else None
        cell.number_format = "General"

    gst = data["gst"]
    sheet["H28"] = f"Rs. {gst['taxable_value']:,.2f}"
    sheet["H30"] = f"Rs. {gst['igst']:,.2f}"
    sheet["H32"] = f"Rs. {gst['cgst']:,.2f}"
    sheet["H34"] = f"Rs. {gst['sgst']:,.2f}"
    sheet["H35"] = f"Rs. {gst['grand_total']:,.2f}"
    for cell in ("H28", "H30", "H32", "H34", "H35"):
        sheet[cell].number_format = "General"
    # The PDF uses the longer `Rs.` marker in place of the workbook's compact
    # rupee glyph. Keep the grand-total value inside its narrow cell.
    total_font = copy(sheet["H35"].font)
    total_font.sz = 9
    sheet["H35"].font = total_font
    sheet["C37"] = amount_in_words(gst["grand_total"])
    sheet["G16"] = "Rate (Rs.)"
    sheet["H16"] = "Amount (Rs.)"
    sheet.print_area = "A1:H41"

    # The workbook has three blank rows beyond its invoice print area. Remove
    # them in the temporary copy; the PDF table renderer must have matching
    # cell and row-height counts.
    for merged_range in list(sheet.merged_cells.ranges):
        if merged_range.max_row > 41:
            sheet.unmerge_cells(str(merged_range))
    if sheet.max_row > 41:
        sheet.delete_rows(42, sheet.max_row - 41)

    # xlsx2html relies on explicit widths; supply the template's default width
    # for its otherwise-unspecified D column so merged cells stay aligned.
    if not sheet.column_dimensions["D"].customWidth:
        sheet.column_dimensions["D"].width = sheet.sheet_format.defaultColWidth or 13

    # Expand row heights to use the available A4 print area. Column widths stay
    # at their original proportions and the PDF table is fitted to page width.
    for row in range(1, 42):
        dimension = sheet.row_dimensions[row]
        dimension.height = (dimension.height or sheet.sheet_format.defaultRowHeight or 15) * 1.13
    for row in range(42, sheet.max_row + 1):
        sheet.row_dimensions[row].hidden = True
    workbook.save(destination)


def _no_extra_rows_or_columns(*_args):
    return None


def convert_xlsx_to_pdf(
    xlsx_path: str | Path, pdf_path: str | Path, data: dict
) -> str:
    """Convert the invoice XLSX itself to a styled A4 PDF without Office."""
    xlsx_path = Path(xlsx_path).resolve()
    pdf_path = Path(pdf_path).resolve()
    if not xlsx_path.is_file():
        raise FileNotFoundError(f"Invoice workbook not found: {xlsx_path}")
    pdf_path.parent.mkdir(parents=True, exist_ok=True)

    with tempfile.TemporaryDirectory(prefix="invoice-render-") as temp_dir:
        render_copy = Path(temp_dir) / xlsx_path.name
        _prepare_pdf_copy(xlsx_path, data, render_copy)
        html_stream = xlsx2html(
            str(render_copy),
            sheet="Tax Invoice",
            locale="en_IN",
            append_headers=_no_extra_rows_or_columns,
            append_lineno=_no_extra_rows_or_columns,
        )
        html = html_stream.getvalue()
        # xlsx2html emits Excel point sizes as CSS px values. CSS pixels render
        # at 0.75pt in PDF, so convert the numeric size back to pt to preserve
        # the font size defined by the workbook. Excel's font-weight and cell
        # alignment are already present in the generated cell styles.
        html = re.sub(
            r"font-size:\s*([\d.]+)px",
            lambda match: f"font-size: {match.group(1)}pt",
            html,
        )
        # Excel leaves a small inset between cell text and its grid lines.
        # Add that inset to the HTML table and a little extra before the final
        # amount, whose `Rs.` prefix is wider than Excel's rupee symbol.
        html = html.replace(
            'id="Tax Invoice!F35" style="',
            'id="Tax Invoice!F35" style="padding-right: 6px;',
        )
        page_style = """
        <style>
          @page { size: A4 portrait; margin: 8mm; }
          html, body { margin: 0; padding: 0; }
          table, td { font-family: Helvetica, Arial, sans-serif; }
          td { padding: 1px 2px; }
          table { width: 100%; border-collapse: collapse;
                  -pdf-keep-in-frame-mode: shrink; }
        </style>
        """
        html = html.replace("</head>", page_style + "</head>", 1)
        with pdf_path.open("wb") as output:
            result = pisa.CreatePDF(html, dest=output, encoding="utf-8")
        if result.err:
            raise RuntimeError("Python could not render the invoice workbook as PDF.")

    if not pdf_path.is_file() or pdf_path.stat().st_size == 0:
        raise RuntimeError("PDF renderer finished without creating an invoice PDF.")
    return str(pdf_path)

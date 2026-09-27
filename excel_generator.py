"""Fill a copy of the supplied Hanuman invoice workbook."""

from pathlib import Path

from openpyxl import load_workbook


BASE_DIR = Path(__file__).resolve().parent
TEMPLATE_PATH = BASE_DIR / "Hanuman.xlsx"


def generate_invoice_xlsx(data: dict, out_path: str | Path) -> str:
    """Create an editable invoice workbook while preserving the template layout."""
    if not TEMPLATE_PATH.is_file():
        raise FileNotFoundError(f"Invoice template not found: {TEMPLATE_PATH}")

    items = data["items"]
    if not items:
        raise ValueError("An invoice must contain at least one item.")
    if len(items) > 11:
        raise ValueError("The supplied invoice template supports up to 11 items.")

    workbook = load_workbook(TEMPLATE_PATH)
    sheet = workbook["Tax Invoice"]

    # Keep the invoice readable and centered when Excel exports it to PDF.
    sheet.print_area = "A1:H41"
    sheet.page_setup.orientation = "portrait"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_options.horizontalCentered = True
    sheet.print_options.verticalCentered = True
    sheet.page_margins.left = 0.25
    sheet.page_margins.right = 0.25
    sheet.page_margins.top = 0.3
    sheet.page_margins.bottom = 0.3
    sheet.page_margins.header = 0.1
    sheet.page_margins.footer = 0.1

    # Invoice metadata in the existing template.
    values = {
        "G4": data["invoice_no"],
        "G5": data["invoice_date"],
        "G6": data.get("delivery_note") or "-",
        "G7": data.get("payment_terms") or "-",
        "G8": data.get("vehicle_no") or "-",
        "G10": data.get("buyer_order_no") or "-",
        "G11": data.get("dispatch_doc_no") or "-",
        "A11": data["buyer_name"],
        "A14": f"GSTIN/UIN: {data.get('buyer_gstin') or '-'}",
        "G12": data.get("destination") or "-",
        "G13": data.get("delivery_terms") or "As agreed",
    }
    for address_cell in ("A12", "A13"):
        sheet[address_cell] = None
    address_lines = [line.strip() for line in data["buyer_address"].splitlines() if line.strip()]
    if address_lines:
        sheet["A12"] = address_lines[0]
        if len(address_lines) > 1:
            sheet["A13"] = ", ".join(address_lines[1:])
    else:
        sheet["A12"] = data["buyer_address"].strip()
    for cell, value in values.items():
        sheet[cell] = value

    # The existing invoice template has 11 prepared item rows, rows 17 through 27.
    for row in range(17, 28):
        sheet.cell(row, 1).value = None
        sheet.cell(row, 2).value = None
        sheet.cell(row, 5).value = None
        sheet.cell(row, 6).value = None
        sheet.cell(row, 7).value = None
    for index, item in enumerate(items, start=1):
        row = 16 + index
        sheet.cell(row, 1).value = index
        sheet.cell(row, 2).value = item["description"]
        sheet.cell(row, 5).value = item["hsn"]
        sheet.cell(row, 6).value = item["qty"]
        sheet.cell(row, 7).value = item["rate"]
        # Keep the template's blank-safe line amount formula.
        sheet.cell(row, 8).value = f'=IF(AND(F{row}="",G{row}=""),"",F{row}*G{row})'

    # Clear unused item rows while retaining formulas and formatting.
    for row in range(17 + len(items), 28):
        sheet.cell(row, 8).value = f'=IF(AND(F{row}="",G{row}=""),"",F{row}*G{row})'

    gst = data["gst"]
    if gst["gst_type"] == "IGST":
        sheet["H29"] = gst["gst_rate"] / 100
        sheet["H31"] = 0
        sheet["H33"] = 0
    else:
        sheet["H29"] = 0
        sheet["H31"] = gst["gst_rate"] / 200
        sheet["H33"] = gst["gst_rate"] / 200

    # Ask Excel/compatible spreadsheet software to recalculate on open.
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"

    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(out_path)
    return str(out_path)
